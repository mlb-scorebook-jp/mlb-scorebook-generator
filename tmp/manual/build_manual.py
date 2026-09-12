from pathlib import Path
from docx import Document
from docx.enum.section import WD_SECTION
from docx.enum.table import WD_CELL_VERTICAL_ALIGNMENT, WD_TABLE_ALIGNMENT
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Inches, Pt, RGBColor

ROOT = Path("/Users/hiramotoakihiro/Documents/GitHub/mlb-scorebook-generator")
OUT = ROOT / "output" / "manual"
SHOTS = OUT / "screenshots"
DOCX = OUT / "MLB_SCOREBOOK_GENERATOR_User_Manual_JA.docx"

doc = Document()
section = doc.sections[0]
section.page_width = Inches(8.5)
section.page_height = Inches(11)
section.top_margin = Inches(0.65)
section.bottom_margin = Inches(0.65)
section.left_margin = Inches(0.7)
section.right_margin = Inches(0.7)

styles = doc.styles
for style in styles:
    if not hasattr(style, "font"):
        continue
    style.font.name = "Arial Unicode MS"
    style._element.get_or_add_rPr().get_or_add_rFonts().set(qn("w:eastAsia"), "Arial Unicode MS")
for name in ["Normal", "Title", "Subtitle", "Heading 1", "Heading 2", "Heading 3"]:
    style = styles[name]
    style.font.color.rgb = RGBColor(0, 0, 0)

styles["Normal"].font.size = Pt(10.5)
styles["Normal"].paragraph_format.space_after = Pt(6)
styles["Normal"].paragraph_format.line_spacing = 1.18
styles["Title"].font.size = Pt(28)
styles["Title"].font.bold = True
styles["Title"].paragraph_format.space_after = Pt(14)
title_p_pr = styles["Title"]._element.get_or_add_pPr()
title_border = title_p_pr.find(qn("w:pBdr"))
if title_border is not None:
    title_p_pr.remove(title_border)
styles["Subtitle"].font.size = Pt(14)
styles["Subtitle"].font.color.rgb = RGBColor(0, 0, 0)
styles["Heading 1"].font.size = Pt(20)
styles["Heading 1"].font.bold = True
styles["Heading 1"].paragraph_format.space_before = Pt(14)
styles["Heading 1"].paragraph_format.space_after = Pt(8)
styles["Heading 2"].font.size = Pt(14)
styles["Heading 2"].font.bold = True
styles["Heading 2"].paragraph_format.space_before = Pt(10)
styles["Heading 2"].paragraph_format.space_after = Pt(5)
styles["Heading 3"].font.size = Pt(11)
styles["Heading 3"].font.bold = True

def set_cell_shading(cell, fill):
    tc_pr = cell._tc.get_or_add_tcPr()
    shd = OxmlElement("w:shd")
    shd.set(qn("w:fill"), fill)
    tc_pr.append(shd)

def set_cell_margins(cell, top=100, start=110, bottom=100, end=110):
    tc = cell._tc
    tc_pr = tc.get_or_add_tcPr()
    tc_mar = tc_pr.first_child_found_in("w:tcMar")
    if tc_mar is None:
        tc_mar = OxmlElement("w:tcMar")
        tc_pr.append(tc_mar)
    for margin, value in [("top", top), ("start", start), ("bottom", bottom), ("end", end)]:
        node = tc_mar.find(qn(f"w:{margin}"))
        if node is None:
            node = OxmlElement(f"w:{margin}")
            tc_mar.append(node)
        node.set(qn("w:w"), str(value))
        node.set(qn("w:type"), "dxa")

def set_repeat_table_header(row):
    tr_pr = row._tr.get_or_add_trPr()
    tbl_header = OxmlElement("w:tblHeader")
    tbl_header.set(qn("w:val"), "true")
    tr_pr.append(tbl_header)

def add_page_number(paragraph):
    paragraph.alignment = WD_ALIGN_PARAGRAPH.CENTER
    run = paragraph.add_run()
    fld_char1 = OxmlElement("w:fldChar")
    fld_char1.set(qn("w:fldCharType"), "begin")
    instr = OxmlElement("w:instrText")
    instr.set(qn("xml:space"), "preserve")
    instr.text = "PAGE"
    fld_char2 = OxmlElement("w:fldChar")
    fld_char2.set(qn("w:fldCharType"), "end")
    run._r.extend([fld_char1, instr, fld_char2])
    run.font.size = Pt(9)

header = section.header.paragraphs[0]
header.text = "MLB SCOREBOOK GENERATOR 取扱説明書"
header.alignment = WD_ALIGN_PARAGRAPH.RIGHT
header.runs[0].font.size = Pt(8)
header.runs[0].font.color.rgb = RGBColor(90, 90, 90)
add_page_number(section.footer.paragraphs[0])

def new_page_heading(text):
    p = doc.add_heading(text, level=1)
    p.paragraph_format.page_break_before = True
    return p

def add_bullet(text, level=0):
    p = doc.add_paragraph(style="List Bullet" if level == 0 else "List Bullet 2")
    p.add_run(text)
    return p

def add_step(number, title, body):
    p = doc.add_paragraph()
    p.paragraph_format.space_before = Pt(7)
    r = p.add_run(f"{number}  {title}")
    r.bold = True
    r.font.size = Pt(11.5)
    body_p = doc.add_paragraph(body)
    body_p.paragraph_format.left_indent = Inches(0.25)

def add_image(name, caption, width=7.0):
    path = SHOTS / name
    p = doc.add_paragraph()
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    p.paragraph_format.keep_with_next = True
    p.add_run().add_picture(str(path), width=Inches(width))
    cap = doc.add_paragraph(caption)
    cap.alignment = WD_ALIGN_PARAGRAPH.CENTER
    cap.paragraph_format.space_after = Pt(9)
    cap.runs[0].italic = True
    cap.runs[0].font.size = Pt(9)
    cap.runs[0].font.color.rgb = RGBColor(80, 80, 80)

def add_table(headers, rows, widths):
    table = doc.add_table(rows=1, cols=len(headers))
    table.alignment = WD_TABLE_ALIGNMENT.CENTER
    table.style = "Table Grid"
    hdr = table.rows[0]
    set_repeat_table_header(hdr)
    for i, text in enumerate(headers):
        hdr.cells[i].text = text
        hdr.cells[i].width = Inches(widths[i])
        hdr.cells[i].vertical_alignment = WD_CELL_VERTICAL_ALIGNMENT.CENTER
        set_cell_shading(hdr.cells[i], "243746")
        set_cell_margins(hdr.cells[i])
        for run in hdr.cells[i].paragraphs[0].runs:
            run.font.bold = True
            run.font.color.rgb = RGBColor(255, 255, 255)
            run.font.size = Pt(9.5)
    for row_index, values in enumerate(rows):
        cells = table.add_row().cells
        for i, value in enumerate(values):
            cells[i].text = value
            cells[i].width = Inches(widths[i])
            cells[i].vertical_alignment = WD_CELL_VERTICAL_ALIGNMENT.CENTER
            set_cell_margins(cells[i])
            if row_index % 2:
                set_cell_shading(cells[i], "F2F6F8")
            for run in cells[i].paragraphs[0].runs:
                run.font.size = Pt(9.5)
    return table

# Cover
p = doc.add_paragraph(style="Title")
p.alignment = WD_ALIGN_PARAGRAPH.CENTER
p.add_run("MLB SCOREBOOK GENERATOR")
p_pr = p._p.get_or_add_pPr()
p_bdr = p_pr.find(qn("w:pBdr"))
if p_bdr is not None:
    p_pr.remove(p_bdr)
p2 = doc.add_paragraph()
p2.alignment = WD_ALIGN_PARAGRAPH.CENTER
p2.paragraph_format.space_after = Pt(12)
subtitle_run = p2.add_run("取扱説明書")
subtitle_run.bold = True
subtitle_run.font.size = Pt(16)
intro = doc.add_paragraph(
    "MLB公式データから試合情報を読み込み、スコアブック、投球チャート、試合前資料、"
    "当日の記録候補を確認するための操作手順をまとめています。初めて使う場合は、"
    "次ページの基本操作から順に進めてください。"
)
intro.alignment = WD_ALIGN_PARAGRAPH.CENTER
intro.paragraph_format.space_after = Pt(12)
add_image("01-home.png", "起動画面")
meta = doc.add_paragraph("Version 2.3.8　Updated 2026-09-12")
meta.alignment = WD_ALIGN_PARAGRAPH.CENTER

new_page_heading("最初に確認すること")
doc.add_paragraph(
    "このアプリはMLB公式APIからデータを取得します。試合を開くときはインターネット接続が必要です。"
    "表示内容は試合データの更新状況によって変わるため、放送や公開資料に使う情報は、"
    "画面内のGamedayや公式記事リンクでも確認してください。"
)
doc.add_heading("基本の使い分け", level=2)
add_table(
    ["画面", "できること", "主な使い方"],
    [
        ["スコアブック生成", "Game PKからスコアブックを作成", "試合結果の整理、PDFやPNGの出力"],
        ["試合前情報", "先発、対戦成績、注目選手、球場履歴などを表示", "中継前の資料確認"],
        ["本日の記録", "記録や特殊事象を自動検出", "話題候補と根拠リンクの確認"],
        ["投球チャート", "球種、コース、打者別傾向を表示", "投手内容の確認"]
    ],
    [1.35, 2.7, 3.0]
)
doc.add_heading("画面上部の共通メニュー", level=2)
add_bullet("スコアブック生成　試合を読み込んでスコアを表示します。")
add_bullet("試合前情報　対象日の全試合と選手情報を表示します。")
add_bullet("本日の記録　対象日の記録候補を再集計します。")
add_bullet("MLB SCOREBOOK GENERATORのタイトル　トップ画面へ戻ります。")

new_page_heading("スコアブックを作る")
add_step(1, "スコアブック生成を開く", "トップ画面または上部メニューからスコアブック生成を選びます。")
add_step(2, "試合を選ぶ", "本日の試合一覧から選ぶか、URL Game PKを直接入力を開きます。")
add_step(3, "Game PKを入力する", "MLB GamedayのURL、またはURL内の数字のGame PKを入力します。")
add_step(4, "読み込みを開始する", "保存を押します。試合終了後のデータは通常数秒で表示されます。")
doc.add_heading("Game PKとは", level=2)
doc.add_paragraph(
    "MLBが各試合に付けている数字の識別番号です。MLB GamedayのURLをそのまま貼り付けても、"
    "アプリがGame PKを読み取ります。数字だけを入力する場合は、別の試合番号を入れないよう確認してください。"
)
doc.add_heading("試合が一覧に見つからない場合", level=2)
add_bullet("日付が日本時間ではなく現地日付になっているか確認します。")
add_bullet("ダブルヘッダーは第1試合と第2試合でGame PKが異なります。")
add_bullet("中止や延期の直後はMLB公式データの反映を待ってから再読み込みします。")

new_page_heading("スコアブック画面")
add_image("05-scorebook.png", "スコアブックの表示例", 5.7)
doc.add_paragraph(
    "左側に試合情報と設定、中央に先攻と後攻のスコアブック、画面下部に出力ボタンがあります。"
    "横幅が足りない場合は、下部の横スクロールを使って右側の回や投手欄を確認します。"
)
add_table(
    ["場所", "内容"],
    [
        ["上段", "イニング別得点、R H E、対戦成績、日付、球場、天候、試合時間"],
        ["打者欄", "打順、背番号、選手名、守備位置、各打席結果、打数、安打、得点、打点など"],
        ["投手欄", "球数、投球回、被安打、失点、自責点、四死球、奪三振、防御率"],
        ["特記事項", "選手交代、守備変更、チャレンジなどの補足情報"],
        ["画面下部", "PDF、PNG、PRINT、JSONの出力"]
    ],
    [1.35, 5.7]
)

new_page_heading("打席結果と公式PBPを確認する")
doc.add_paragraph(
    "スコアブック内の打席結果をクリックすると、その打席に対応するMLB公式ページを開けます。"
    "左側の設定でリンク先を選択してください。"
)
doc.add_heading("MLB Gameday", level=2)
doc.add_paragraph(
    "MLB公式Gamedayの該当プレーを開きます。プレー説明、投球結果、走者の動きなどを確認する場合に使います。"
)
doc.add_heading("Research Tool 映像付き", level=2)
doc.add_paragraph(
    "MLB Research Toolの映像付きプレーを開きます。対象プレーに映像がない場合や、"
    "MLB側の仕様変更があった場合は、Gamedayで確認してください。"
)
doc.add_heading("主な表記の見方", level=2)
add_table(
    ["例", "意味"],
    [
        ["K または Ks", "三振。振り逃げなどは周囲の記号も確認します。"],
        ["BB", "四球"],
        ["HBP", "死球"],
        ["HR", "本塁打"],
        ["Eと守備番号", "失策"],
        ["PH PR", "代打、代走"],
        ["守備位置の矢印", "試合中の守備位置変更"]
    ],
    [1.6, 5.45]
)
doc.add_paragraph(
    "表記だけで判断しにくいプレーは、打席結果から公式PBPを開いて説明文と走者情報を確認してください。"
)

new_page_heading("スコアブックを出力する")
add_table(
    ["ボタン", "出力内容", "用途"],
    [
        ["PDF", "印刷用PDFを作成", "共有、保存、紙資料"],
        ["PNG", "スコアブックを画像化", "画像共有、資料への貼り付け"],
        ["PRINT", "ブラウザの印刷画面を開く", "直接印刷、プリンタ設定"],
        ["JSON", "読み込んだ試合状態を保存", "バックアップ、後日の再利用"]
    ],
    [1.0, 2.7, 3.35]
)
doc.add_heading("出力前の確認", level=2)
add_bullet("試合状態が試合終了になっていることを確認します。")
add_bullet("延長戦の場合は最終回まで表示されていることを確認します。")
add_bullet("選手交代、投手成績、得点、安打、失策を公式Box Scoreと照合します。")
add_bullet("印刷時に横向き指定が必要な場合は、PRINT後の印刷設定で変更します。")
doc.add_heading("JSONを使う場合", level=2)
doc.add_paragraph(
    "JSONはスコアブックの画像ではなく、読み込んだ状態を保存するファイルです。"
    "試合データを保存しておく場合は、Game PKが分かる名前のまま保管してください。"
)

new_page_heading("投球チャートを使う")
add_image("06-pitch-chart.png", "投球チャートの表示例", 6.2)
doc.add_paragraph(
    "スコアブックの投手成績欄にある投球チャートを押すと、球種別の投球位置、打者別コース、"
    "球種構成を表示します。上部のスコアへ戻るで元のスコアブックへ戻れます。"
)
add_step(1, "投手を選ぶ", "表示されている投手の中から確認したい投手を探します。")
add_step(2, "球種を絞る", "球種の選択欄で4シーム、スライダーなどを指定します。")
add_step(3, "打者を絞る", "打者の選択欄で相手打者を指定し、コース別傾向を確認します。")
add_step(4, "球種構成を比較する", "当日の割合とシーズン割合を比較します。")
doc.add_paragraph(
    "ストライクゾーン表示はMLBのトラッキングデータを使います。判定そのものを示す図ではないため、"
    "際どい投球はGamedayの投球詳細も併せて確認してください。"
)

new_page_heading("試合前情報を開く")
add_image("02-pregame-top.png", "試合前情報トップの表示例", 6.5)
doc.add_paragraph(
    "試合前情報トップでは、対象日の日本人選手、当日成績、全試合、順位やポストシーズン情報を確認できます。"
    "日付欄または左右の矢印で対象日を変更します。"
)
add_step(1, "日付を選ぶ", "現地日付を選びます。今日を押すと現在のMLB日付へ戻ります。")
add_step(2, "対象試合を選ぶ", "全試合のカードから確認したい試合を押します。")
add_step(3, "選手を開く", "日本人選手カードを押すと、その選手の詳細へ移動します。")
doc.add_heading("時差に注意する", level=2)
doc.add_paragraph(
    "日本時間の午前中に行われる試合は、MLB現地日付では前日になることがあります。"
    "放送日ではなく、試合カードに表示される現地日付を基準に選んでください。"
)

new_page_heading("試合前情報の詳細")
add_image("03-pregame-game.png", "先発投手と球場履歴を含む試合詳細", 6.5)
doc.add_paragraph(
    "試合詳細には、対戦カード、順位、先発投手、直近登板、注目選手、チーム動向、"
    "公式記事などが表示されます。数値や過去試合の日付がリンクになっている場合は、"
    "クリックすると根拠ページを新しいタブで開きます。"
)
doc.add_heading("先発投手欄", level=2)
add_bullet("今季成績　登板数、先発数、勝敗、防御率")
add_bullet("対戦成績　相手球団との通算登板、勝敗、防御率")
add_bullet("球場成績　当該球場での通算登板、勝敗、防御率")
add_bullet("直近登板　日付、相手、勝敗、投球回、失点、奪三振、四死球")
doc.add_heading("球場に関する注目情報", level=2)
doc.add_paragraph(
    "過去に同じ球場でノーヒッター、継投ノーヒッター、完封、15奪三振以上、"
    "ポストシーズン登板がある場合は再登板情報を表示します。"
    "過去10登板で勝利がない場合や3連敗以上の場合は、マイナス材料として表示します。"
    "各日付の根拠リンクからMLB公式Gamedayを開けます。"
)

new_page_heading("本日の記録を確認する")
add_image("04-daily-records.png", "本日の記録の表示例", 5.5)
doc.add_paragraph(
    "本日の記録は、対象日のFinal試合を解析し、個人記録、チーム記録、特殊事象、"
    "日本人選手、誕生日などの話題候補を一覧にします。解析完了までは進行状況が表示されます。"
)
add_step(1, "対象日を選ぶ", "現地日付を指定します。")
add_step(2, "再調査を押す", "最新のMLB公式データを使って再解析します。")
add_step(3, "記録を確認する", "記録名、選手、前回記録、通算集計を確認します。")
add_step(4, "根拠を開く", "Gameday、前回のGameday、MLB公式記事などのリンクを開きます。")
doc.add_heading("表示結果の扱い", level=2)
doc.add_paragraph(
    "特殊事象には自動判定による候補が含まれます。ナレーション原稿や公開資料へ使用する前に、"
    "Gameday、Box Score、公式記事で事実関係を確認してください。"
)

new_page_heading("記録検索を使う")
doc.add_paragraph(
    "本日の記録画面で記録検索を選ぶと、過去の記録をキーワード、日付、球団、選手などで探せます。"
    "複数の条件を組み合わせると検索結果を絞り込めます。"
)
add_table(
    ["検索例", "探せる内容"],
    [
        ["大谷 3本塁打", "大谷選手の1試合3本塁打"],
        ["LAD サイクル", "ドジャースに関係するサイクル安打"],
        ["2024-07-22", "指定日の記録"],
        ["イマキュレート", "イマキュレートイニング"],
        ["野手登板", "野手が投手として登板した試合"]
    ],
    [2.2, 4.85]
)
doc.add_heading("検索で見つからない場合", level=2)
add_bullet("選手名をフルネームではなく名字だけで試します。")
add_bullet("球団名を3文字コードで試します。")
add_bullet("日付はYYYY-MM-DD形式で入力します。")
add_bullet("記録名を短くして再検索します。")

new_page_heading("選手交代と特記事項を確認する")
doc.add_paragraph(
    "代打、代走、投手交代、守備交代、守備位置変更はスコアブックと特記事項に反映されます。"
    "複数の変更が同時に行われた場合は、スコアブック上の表記とMLB公式PBPの交代説明を照合してください。"
)
add_table(
    ["項目", "確認点"],
    [
        ["代打", "元の打者と代打選手、打席が行われた回"],
        ["代走", "交代した走者と、その後の進塁や得点"],
        ["投手交代", "交代した回、対戦した最初の打者、投手成績"],
        ["守備交代", "選手名、守備位置、打順"],
        ["守備位置変更", "変更前と変更後の守備位置"],
        ["チャレンジ", "対象プレー、判定結果、残り回数"]
    ],
    [1.55, 5.5]
)
doc.add_paragraph(
    "MLB公式データが試合後に訂正された場合、再読み込みすると表示も更新されます。"
    "保存済みのPDFやPNGは自動更新されないため、必要に応じて再出力してください。"
)

new_page_heading("スマートフォンと印刷")
doc.add_heading("スマートフォンで見る", level=2)
add_bullet("横に長いスコアブックは横スクロールして確認します。")
add_bullet("投球チャートや表は、端末を横向きにすると見やすくなります。")
add_bullet("リンクは新しいタブで開くため、元の画面へ戻るときはブラウザのタブを切り替えます。")
doc.add_heading("印刷する", level=2)
add_step(1, "PRINTを押す", "ブラウザの印刷画面を開きます。")
add_step(2, "用紙方向を選ぶ", "スコアブック全体を印刷する場合は横向きを確認します。")
add_step(3, "倍率を確認する", "右端が切れる場合は用紙に合わせる、または倍率を下げます。")
add_step(4, "プレビューを確認する", "ページ分割、余白、選手名の切れがないことを確認して印刷します。")
doc.add_heading("PDFとPNGの使い分け", level=2)
doc.add_paragraph(
    "紙資料や複数ページの共有にはPDF、原稿やスライドへ一部分を貼る場合はPNGが向いています。"
)

new_page_heading("困ったとき")
add_table(
    ["症状", "確認と対処"],
    [
        ["試合を読み込めない", "インターネット接続、Game PK、MLB公式サイトの状態を確認して再読み込みします。"],
        ["試合一覧に出ない", "現地日付を確認し、Game PKの直接入力を試します。"],
        ["選手名や交代が違う", "試合終了後に再読み込みし、公式Box ScoreとPBPを確認します。"],
        ["本日の記録が途中で止まる", "少し待ってから再調査を押します。対象試合数が多い日は時間がかかります。"],
        ["リンクが開かない", "ブラウザのポップアップ制限とネット接続を確認します。"],
        ["PDFやPNGが保存されない", "ブラウザのダウンロード許可と保存先を確認します。"],
        ["印刷で右端が切れる", "横向き、用紙に合わせる、余白を狭くする設定を試します。"],
        ["表示が古い", "ブラウザを再読み込みします。必要ならキャッシュを更新します。"]
    ],
    [2.05, 5.0]
)
doc.add_heading("確認の優先順位", level=2)
doc.add_paragraph(
    "表示に疑問がある場合は、最初に画面内のGamedayリンクを開き、次にMLB公式Box Score、"
    "公式記事の順で確認します。アプリの自動判定と公式情報が異なる場合は、MLB公式の最新情報を優先してください。"
)

new_page_heading("利用上の注意")
add_bullet("試合中のデータは速報値で、試合後に訂正される場合があります。")
add_bullet("記録候補や特殊事象は自動検出です。放送前に根拠リンクで確認してください。")
add_bullet("MLB公式サイトの画面やURL仕様が変わると、一部リンクが正しく開かない場合があります。")
add_bullet("投球位置や球種はMLBのトラッキングデータに基づきます。公式記録と判定目的を混同しないでください。")
add_bullet("出力したPDF、PNG、JSONには試合情報が含まれます。共有先と保存場所を確認してください。")
doc.add_heading("日常的な確認手順", level=2)
add_step(1, "試合前", "試合前情報で先発、球場履歴、注目選手、公式記事を確認します。")
add_step(2, "試合後", "スコアブックを読み込み、得点、安打、失策、投手成績、交代を確認します。")
add_step(3, "話題確認", "本日の記録で候補を確認し、Gamedayや公式記事で裏付けます。")
add_step(4, "資料保存", "必要な形式でPDF、PNG、JSONを保存します。")

OUT.mkdir(parents=True, exist_ok=True)
for paragraph in list(doc.paragraphs) + [
    paragraph
    for table in doc.tables
    for row in table.rows
    for cell in row.cells
    for paragraph in cell.paragraphs
]:
    for run in paragraph.runs:
        run.font.name = "Arial Unicode MS"
        run._element.get_or_add_rPr().get_or_add_rFonts().set(
            qn("w:eastAsia"), "Arial Unicode MS"
        )
doc.save(DOCX)
print(DOCX)
